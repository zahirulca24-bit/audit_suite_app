import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(r'''
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="AAAAAA"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="AAAAAA"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def generate_annual_return_package(output_dir, company_name, reg_no, reg_address, year_config, shareholders):
    os.makedirs(output_dir, exist_ok=True)
    fy = year_config["fy"]
    agm_no = year_config.get("agm_no", "Annual")
    agm_date = year_config["agm_date"]
    notice_date = year_config.get("notice_date", "14 days prior")
    auditor_name = year_config.get("auditor_name", "Statutory Auditors")
    next_auditor = year_config.get("next_auditor", "Statutory Auditors")
    remuneration = year_config.get("remuneration", "Tk. 25,000/-")

    # 1. AGM Notice
    doc_notice = docx.Document()
    p = doc_notice.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{company_name.upper()}\n").bold = True
    p.add_run(f"{reg_address}\nRegistration No: {reg_no}\n\n")
    p.add_run(f"NOTICE OF THE {agm_no.upper()} ANNUAL GENERAL MEETING\n").bold = True

    p_body = doc_notice.add_paragraph()
    p_body.paragraph_format.line_spacing = 1.15
    p_body.add_run(
        f"Notice is hereby given that the {agm_no} Annual General Meeting of {company_name.upper()} "
        f"will be held at the Registered Office on {agm_date} at 11:00 AM to transact the following business:\n\n"
        f"1. To receive, consider and adopt the Audited Financial Statements for the year ended 30th June {fy.split('-')[-1]} together with Reports of the Directors and Auditors thereon.\n"
        f"2. To re-elect Directors retiring by rotation.\n"
        f"3. To appoint Statutory Auditors ({next_auditor}) and fix their remuneration.\n\n"
        f"Dated: {notice_date}\n\n\n"
        "By Order of the Board,\n\n"
        "_________________________\n"
        "Managing Director"
    )
    notice_file = os.path.join(output_dir, f"AGM_Notice_FY_{fy}.docx")
    doc_notice.save(notice_file)

    # 2. AGM Minutes
    doc_min = docx.Document()
    p_m = doc_min.add_paragraph()
    p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_m.add_run(f"{company_name.upper()}\n").bold = True
    p_m.add_run(f"MINUTES OF THE {agm_no.upper()} ANNUAL GENERAL MEETING\nHELD ON {agm_date.upper()}\n\n").bold = True

    p_mb = doc_min.add_paragraph()
    p_mb.paragraph_format.line_spacing = 1.15
    p_mb.add_run(
        f"The {agm_no} Annual General Meeting of the shareholders of {company_name.upper()} "
        f"was held on {agm_date} at 11:00 AM at the registered office.\n\n"
        "RESOLUTIONS ADOPTED:\n"
        f"1. The Audited Balance Sheet and Income Statement for FY {fy} as audited by {auditor_name} were unanimously approved and adopted.\n"
        f"2. {next_auditor} were appointed as Statutory Auditors for the next financial year with remuneration of {remuneration}.\n\n\n"
        "_________________________\n"
        "Chairman of the Meeting"
    )
    min_file = os.path.join(output_dir, f"AGM_Minutes_FY_{fy}.docx")
    doc_min.save(min_file)

    # 3. Schedule X (Form X)
    doc_x = docx.Document()
    p_x = doc_x.add_paragraph()
    p_x.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_x.add_run("THE COMPANIES ACT, 1994\nSCHEDULE X\nFORM OF ANNUAL RETURN OF A COMPANY HAVING SHARE CAPITAL\n[Section 36]\n\n").bold = True
    p_x.add_run(f"Company: {company_name.upper()} | Reg No: {reg_no}\nAnnual Return made up to {agm_date}\n\n")

    tbl = doc_x.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    hdrs = ["Sl", "Name & Address of Shareholder", "Father / Mother", "No. of Shares"]
    for i, h in enumerate(hdrs):
        r = tbl.rows[0].cells[i].paragraphs[0].add_run(h)
        r.bold = True

    for idx, s in enumerate(shareholders, 1):
        rw = tbl.add_row().cells
        rw[0].paragraphs[0].add_run(str(idx))
        rw[1].paragraphs[0].add_run(f"{s['name']}\n{s['address']}")
        rw[2].paragraphs[0].add_run(f"F: {s.get('father_name', '')}\nM: {s.get('mother_name', '')}")
        rw[3].paragraphs[0].add_run(f"{s.get('share_count', 0):,}")

    p_sig = doc_x.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(20)
    p_sig.add_run("Certified to be a true copy submitted to RJSC:\n\n\n_________________________\nManaging Director / Director")
    x_file = os.path.join(output_dir, f"Schedule_X_Annual_Return_FY_{fy}.docx")
    doc_x.save(x_file)

    return {
        "notice": notice_file,
        "minutes": min_file,
        "schedule_x": x_file
    }
