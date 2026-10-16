import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def generate_signing_guide(output_path, company_name, assistant_name="Hemadry", items_guide=None):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{company_name.upper()}\n")
    r.bold = True
    r.font.size = Pt(15)

    r_sub = p.add_run(f"EXACT SIGNING, STAMPING & RJSC SUBMISSION GUIDE (FOR {assistant_name.upper()})\n")
    r_sub.bold = True
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = RGBColor(180, 20, 20)

    doc.add_paragraph("_________________________________________________________________________________").alignment = WD_ALIGN_PARAGRAPH.CENTER

    if not items_guide:
        items_guide = [
            ("Form 117 (Share Transfer Deed)", "Transferor, Transferee, 2 Witnesses, and Managing Director verify."),
            ("Board Resolution (Share Transfer)", "All Directors present at the meeting + Certified True Copy by MD."),
            ("Affidavit (Tk. 300 Stamp)", "Print on Tk. 300 Non-Judicial Stamp. Signed by Transferor and Notarized by Notary Public Dhaka."),
            ("Form IX & XII (Directorship)", "Form IX signed by incoming Director; Form XII signed by Managing Director."),
            ("Share Certificates", "Signed by 2 Directors / MD & Chairman with official seal."),
            ("Annual Return (Schedule X)", "Signed by Managing Director / Director.")
        ]

    for idx, (doc_name, instruction) in enumerate(items_guide, 1):
        p_item = doc.add_paragraph()
        p_item.paragraph_format.line_spacing = 1.15
        r_num = p_item.add_run(f"{idx}. {doc_name}:\n")
        r_num.bold = True
        r_num.font.size = Pt(11)
        r_inst = p_item.add_run(f"   👉 নির্দেশিকা: {instruction}\n")
        r_inst.font.size = Pt(10)

    p_bottom = doc.add_paragraph()
    p_bottom.paragraph_format.space_before = Pt(20)
    p_bottom.add_run("⚠️ বিশেষ সতর্কতা:\n"
                     "১. কোনো সই কালো কালির বলপেন ছাড়া অন্য কিছু দিয়ে করাবেন না।\n"
                     "২. এফিডেভিট অবশ্যই নোটারি রেজিস্ট্রিতে এন্ট্রি করিয়ে নোটারি সীল ও স্ট্যাম্প নেবেন।\n"
                     "৩. সব সাইন সম্পন্ন হওয়ার পর স্ক্যান কপি অ্যাপ/হোয়াটসঅ্যাপে পাঠিয়ে অরিজিনাল আরজেএসসি-তে জমা দেবেন।\n\n"
                     "প্রস্তুতকারী: অডিট অ্যান্ড কর্পোরেট লিগ্যাল টিম").bold = True

    doc.save(output_path)
    return output_path
