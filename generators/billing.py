import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="222222"/>
            <w:left w:val="single" w:sz="8" w:space="0" w:color="222222"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="222222"/>
            <w:right w:val="single" w:sz="8" w:space="0" w:color="222222"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_cell_shading(cell, color_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def generate_billing_requisition(output_path, client_name, client_address, invoice_no, bill_date, items, advance_paid=0):
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Header / Firm Name
    p_firm = doc.add_paragraph()
    p_firm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_firm = p_firm.add_run("FAMES & R\n")
    r_firm.bold = True
    r_firm.font.size = Pt(18)
    r_firm.font.color.rgb = RGBColor(18, 52, 86)

    r_sub = p_firm.add_run("Chartered Accountants | Corporate Legal & Tax Consultants\n"
                           "Hossain Tower (11th Floor), 116 Naya Paltan, Box Culvert Road, Dhaka-1000\n"
                           "Email: info@famesr.com | Web: www.famesr.com\n")
    r_sub.font.size = Pt(9.5)
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph("_________________________________________________________________________________").alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("PROFESSIONAL BILL / REQUISITION")
    r_t.bold = True
    r_t.font.size = Pt(13)

    # Metadata
    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Invoice / Ref No: {invoice_no}\n").bold = True
    p_meta.add_run(f"Date: {bill_date}\n\n")
    p_meta.add_run("To:\n").bold = True
    p_meta.add_run(f"The Board of Directors\n{client_name}\n{client_address}\n")

    # Table of Items
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    headers = ["Sl", "Particulars / Description of Services", "Category", "Amount (BDT)"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hdr_cells[i].paragraphs[0].add_run(h)
        run.bold = True
        set_cell_shading(hdr_cells[i], "EAEAEA")
        set_cell_margins(hdr_cells[i], 120, 120, 100, 100)

    total_amount = 0
    for idx, item in enumerate(items, 1):
        row = table.add_row().cells
        row[0].paragraphs[0].add_run(str(idx)).alignment = WD_ALIGN_PARAGRAPH.CENTER
        row[1].paragraphs[0].add_run(item["description"])
        row[2].paragraphs[0].add_run(item.get("category", "Professional Fee"))
        amount = item.get("amount", 0)
        total_amount += amount
        row[3].paragraphs[0].add_run(f"{amount:,}/-").alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for c in row:
            set_cell_margins(c, 80, 80, 100, 100)

    # Total Row
    row_tot = table.add_row().cells
    row_tot[0].paragraphs[0].add_run("")
    r_t_lbl = row_tot[1].paragraphs[0].add_run("TOTAL GROSS BILL")
    r_t_lbl.bold = True
    row_tot[2].paragraphs[0].add_run("")
    r_t_val = row_tot[3].paragraphs[0].add_run(f"{total_amount:,}/-")
    r_t_val.bold = True
    set_cell_shading(row_tot[1], "F5F5F5")
    set_cell_shading(row_tot[3], "F5F5F5")

    # Net payable if advance
    net_payable = total_amount - advance_paid
    if advance_paid > 0:
        row_adv = table.add_row().cells
        row_adv[1].paragraphs[0].add_run("Less: Advance Received")
        row_adv[3].paragraphs[0].add_run(f"({advance_paid:,}/-)")
        
        row_net = table.add_row().cells
        row_net[1].paragraphs[0].add_run("NET PAYABLE AMOUNT").bold = True
        row_net[3].paragraphs[0].add_run(f"{net_payable:,}/-").bold = True
        set_cell_shading(row_net[1], "DDEEFF")
        set_cell_shading(row_net[3], "DDEEFF")

    p_terms = doc.add_paragraph()
    p_terms.paragraph_format.space_before = Pt(14)
    p_terms.add_run("Terms of Payment:\n").bold = True
    p_terms.add_run("• Payment may kindly be made by Cheque/Pay Order in favour of 'FAMES & R'.\n"
                    "• Government fees, non-judicial stamp duties, and notary costs are included as specified.\n\n\n\n\n")

    p_sig = doc.add_paragraph()
    p_sig.add_run("_______________________________\n"
                  "Authorized Signatory / Partner\n"
                  "FAMES & R, Chartered Accountants")

    doc.save(output_path)
    return output_path
